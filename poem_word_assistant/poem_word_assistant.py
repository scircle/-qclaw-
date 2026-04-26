
   
#!/usr/bin/env python3
import argparse
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pypinyin import lazy_pinyin, Style


FIXED_HEADER = "跑团少儿晨读会"

HEADER_FONT = "STHupo"
HEADER_SIZE_PT = 28

POEM_FONT = "KaiTi"

TITLE_SIZE_PT = 20
AUTHOR_SIZE_PT = 20
BODY_SIZE_PT = 16
RUBY_SIZE_PT = 12

END_PUNCT = set("。！？；…")

ALLOWED_DYNASTIES = {
    "先秦", "秦", "汉", "东汉", "魏", "晋", "西晋", "东晋",
    "南北朝", "隋", "唐", "五代", "宋", "北宋", "南宋",
    "辽", "金", "元", "明", "清", "近代"
}

SUSPICIOUS_PATTERNS = [
    r"作者",
    r"朝代",
    r"内容如下",
    r"这首诗",
    r"以下是",
    r"赏析",
    r"注释",
    r"译文",
    r"简介",
]

ALLOWED_PUNCT = set("，。！？；：、“”‘’（）《》〈〉—…、")


def has_chinese(text: str) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in text)


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", "", text.strip())


def is_chinese_char(ch: str) -> bool:
    return "\u4e00" <= ch <= "\u9fff"


def normalize_poem_line_punctuation(line: str) -> str:
    s = line.strip()
    if not s:
        return s
    if s[-1] not in END_PUNCT:
        s += "。"
    return s


def validate_title(title: str):
    if not isinstance(title, str) or not title.strip():
        raise ValueError("title must be a non-empty string")
    t = title.strip()
    if len(t) > 30:
        raise ValueError(f"title too long: {t}")
    if not has_chinese(t):
        raise ValueError(f"title must contain Chinese: {t}")
    for pat in SUSPICIOUS_PATTERNS:
        if re.search(pat, t):
            raise ValueError(f"suspicious title: {t}")


def validate_dynasty(dynasty: str):
    if not isinstance(dynasty, str) or not dynasty.strip():
        raise ValueError("dynasty must be a non-empty string")
    d = dynasty.strip()
    if len(d) > 10:
        raise ValueError(f"dynasty too long: {d}")
    if d not in ALLOWED_DYNASTIES:
        if not has_chinese(d) or len(d) > 6:
            raise ValueError(f"suspicious dynasty: {d}")


def validate_author(author: str):
    if not isinstance(author, str) or not author.strip():
        raise ValueError("author must be a non-empty string")
    a = author.strip()
    if len(a) > 20:
        raise ValueError(f"author too long: {a}")
    if not has_chinese(a):
        raise ValueError(f"author must contain Chinese: {a}")
    for pat in SUSPICIOUS_PATTERNS:
        if re.search(pat, a):
            raise ValueError(f"suspicious author: {a}")


def validate_line(line: str):
    if not isinstance(line, str) or not line.strip():
        raise ValueError("poem line must be non-empty string")
    s = line.strip()
    if len(s) > 50:
        raise ValueError(f"line too long: {s}")
    if not has_chinese(s):
        raise ValueError(f"line must contain Chinese: {s}")
    for pat in SUSPICIOUS_PATTERNS:
        if re.search(pat, s):
            raise ValueError(f"suspicious explanatory text in line: {s}")
    for ch in s:
        if ch.isspace():
            continue
        if "\u4e00" <= ch <= "\u9fff":
            continue
        if ch in ALLOWED_PUNCT:
            continue
        raise ValueError(f"unsupported character in line: {s}")


def validate_lines(lines):
    if not isinstance(lines, list) or len(lines) < 2:
        raise ValueError("lines must contain at least 2 lines")
    if len(lines) > 30:
        raise ValueError("too many lines in one poem")

    normalized_lines = [normalize_poem_line_punctuation(x) for x in lines if x.strip()]
    norm = [normalize_text(x) for x in normalized_lines]

    if len(set(norm)) != len(norm):
        raise ValueError("duplicated poem lines")

    for line in normalized_lines:
        validate_line(line)

    return normalized_lines


def validate_poems(poems):
    if not isinstance(poems, list):
        raise ValueError("poems must be a JSON array")
    if len(poems) < 2:
        raise ValueError("must contain at least 2 poems")

    seen = set()
    normalized = []

    for idx, poem in enumerate(poems, start=1):
        if not isinstance(poem, dict):
            raise ValueError(f"poem #{idx} must be object")

        for key in ("title", "dynasty", "author", "lines"):
            if key not in poem:
                raise ValueError(f"poem #{idx} missing key: {key}")

        validate_title(poem["title"])
        validate_dynasty(poem["dynasty"])
        validate_author(poem["author"])
        normalized_lines = validate_lines(poem["lines"])

        title = poem["title"].strip()
        title_key = normalize_text(title)

        if title_key in seen:
            raise ValueError(f"duplicated title: {title}")

        seen.add(title_key)

        normalized.append({
            "title": title,
            "dynasty": poem["dynasty"].strip(),
            "author": poem["author"].strip(),
            "lines": normalized_lines,
        })

    return normalized


def set_run_font(run, font_name: str, size_pt: int):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)


def add_center_text_paragraph(doc: Document, text: str, font_name: str, size_pt: int):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, font_name, size_pt)
    return p


def append_plain_run(paragraph, text: str, base_size_pt: int):
    r = paragraph.add_run(text)
    set_run_font(r, POEM_FONT, base_size_pt)
    return r


def create_ruby_element(base_char: str, ruby_text: str, base_size_pt: int):
    ruby = OxmlElement("w:ruby")

    ruby_pr = OxmlElement("w:rubyPr")

    ruby_align = OxmlElement("w:rubyAlign")
    ruby_align.set(qn("w:val"), "center")
    ruby_pr.append(ruby_align)

    hps = OxmlElement("w:hps")
    hps.set(qn("w:val"), str(int(RUBY_SIZE_PT * 2)))
    ruby_pr.append(hps)

    hps_raise = OxmlElement("w:hpsRaise")
    hps_raise.set(qn("w:val"), str(int((base_size_pt + 2) * 2)))
    ruby_pr.append(hps_raise)

    hps_base = OxmlElement("w:hpsBaseText")
    hps_base.set(qn("w:val"), str(int(base_size_pt * 2)))
    ruby_pr.append(hps_base)

    lid = OxmlElement("w:lid")
    lid.set(qn("w:val"), "zh-CN")
    ruby_pr.append(lid)

    ruby.append(ruby_pr)

    rt = OxmlElement("w:rt")
    rt_r = OxmlElement("w:r")
    rt_rpr = OxmlElement("w:rPr")

    rt_fonts = OxmlElement("w:rFonts")
    rt_fonts.set(qn("w:ascii"), POEM_FONT)
    rt_fonts.set(qn("w:hAnsi"), POEM_FONT)
    rt_fonts.set(qn("w:eastAsia"), POEM_FONT)
    rt_fonts.set(qn("w:hint"), "eastAsia")
    rt_rpr.append(rt_fonts)

    rt_sz = OxmlElement("w:sz")
    rt_sz.set(qn("w:val"), str(int(RUBY_SIZE_PT * 2)))
    rt_rpr.append(rt_sz)

    rt_szcs = OxmlElement("w:szCs")
    rt_szcs.set(qn("w:val"), str(int(RUBY_SIZE_PT * 2)))
    rt_rpr.append(rt_szcs)

    rt_r.append(rt_rpr)

    rt_t = OxmlElement("w:t")
    rt_t.text = ruby_text
    rt_r.append(rt_t)
    rt.append(rt_r)
    ruby.append(rt)

    rb = OxmlElement("w:rubyBase")
    rb_r = OxmlElement("w:r")
    rb_rpr = OxmlElement("w:rPr")

    rb_fonts = OxmlElement("w:rFonts")
    rb_fonts.set(qn("w:ascii"), POEM_FONT)
    rb_fonts.set(qn("w:hAnsi"), POEM_FONT)
    rb_fonts.set(qn("w:eastAsia"), POEM_FONT)
    rb_fonts.set(qn("w:hint"), "eastAsia")
    rb_rpr.append(rb_fonts)

    rb_sz = OxmlElement("w:sz")
    rb_sz.set(qn("w:val"), str(int(base_size_pt * 2)))
    rb_rpr.append(rb_sz)

    rb_szcs = OxmlElement("w:szCs")
    rb_szcs.set(qn("w:val"), str(int(base_size_pt * 2)))
    rb_rpr.append(rb_szcs)

    rb_r.append(rb_rpr)

    rb_t = OxmlElement("w:t")
    rb_t.text = base_char
    rb_r.append(rb_t)
    rb.append(rb_r)
    ruby.append(rb)

    return ruby


def add_ruby_text_paragraph(doc: Document, text: str, base_size_pt: int):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for ch in text:
        if is_chinese_char(ch):
            py = lazy_pinyin(ch, style=Style.TONE)[0]
            p._p.append(create_ruby_element(ch, py, base_size_pt))
        else:
            append_plain_run(p, ch, base_size_pt)

    return p


def build_doc(poems):
    doc = Document()

    add_center_text_paragraph(doc, FIXED_HEADER, HEADER_FONT, HEADER_SIZE_PT)

    for idx, poem in enumerate(poems):
        add_ruby_text_paragraph(doc, poem["title"], TITLE_SIZE_PT)
        add_ruby_text_paragraph(
            doc,
            f"【{poem['dynasty']}】{poem['author']}",
            AUTHOR_SIZE_PT
        )

        for line in poem["lines"]:
            add_ruby_text_paragraph(doc, line, BODY_SIZE_PT)

        if idx != len(poems) - 1:
            doc.add_paragraph("")
            doc.add_paragraph("")

    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--poems", required=True, help="validated poem JSON array")
    parser.add_argument("--output", required=True, help="output docx path")
    args = parser.parse_args()

    try:
        poems_raw = json.loads(args.poems)
        poems = validate_poems(poems_raw)

        outdir = os.path.dirname(args.output)
        if outdir:
            os.makedirs(outdir, exist_ok=True)

        doc = build_doc(poems)
        doc.save(args.output)

        print(f"POEM_DOC_GENERATED: {args.output}")

    except Exception as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
